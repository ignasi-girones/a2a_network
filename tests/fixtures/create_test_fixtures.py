"""Generate test fixtures for Fase 5 manual e2e verification.

Run:  python tests/fixtures/create_test_fixtures.py

Creates in tests/fixtures/:
  - large_50page.pdf    — 50-page PDF to test MAX_CHARS_PER_FILE truncation
  - multisheet.xlsx     — Excel with 3 sheets to test markdown table extraction
  - simple_1page.pdf    — 1-page PDF for the happy-path baseline

These complement the malicious PDFs from create_malicious_pdf.py.
Upload each via the frontend and check the orchestrator logs for:
  - context_brief appears in normalizer output
  - Truncation logged for the 50-page PDF
  - All 3 sheets appear in the Excel extraction
"""

from pathlib import Path

FIXTURES_DIR = Path(__file__).parent


def create_50page_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = FIXTURES_DIR / "large_50page.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4

    for page_num in range(1, 51):
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, h - 72, f"Annual Report 2025 — Page {page_num}")
        c.setFont("Helvetica", 10)
        y = h - 110
        for para in range(8):
            text = (
                f"Section {page_num}.{para + 1}: Revenue in region "
                f"{chr(65 + para)} was ${page_num * 100 + para * 37}M, "
                f"representing a {5 + para}% year-over-year growth. "
                f"Operating margins improved to {20 + para}% driven by "
                f"efficiency initiatives and headcount optimization. "
                f"Customer retention in this segment reached {90 + para}%."
            )
            # Split long lines to fit on page
            while text:
                c.drawString(72, y, text[:90])
                text = text[90:]
                y -= 14
            y -= 6
        c.showPage()

    c.save()
    print(f"Created: {path}  ({page_num} pages)")


def create_multisheet_xlsx():
    from openpyxl import Workbook

    path = FIXTURES_DIR / "multisheet.xlsx"
    wb = Workbook()

    # Sheet 1: Revenue by quarter
    ws1 = wb.active
    ws1.title = "Revenue"
    ws1.append(["Quarter", "Product A", "Product B", "Product C", "Total"])
    for q in range(1, 5):
        a, b, c = q * 120, q * 85, q * 200
        ws1.append([f"Q{q} 2025", a, b, c, a + b + c])
    ws1.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", "=SUM(E2:E5)"])

    # Sheet 2: Employee data
    ws2 = wb.create_sheet("Employees")
    ws2.append(["Department", "Headcount", "Avg Salary", "Turnover %"])
    for dept, hc, sal, turn in [
        ("Engineering", 120, 95000, 8.2),
        ("Sales", 85, 72000, 14.5),
        ("Marketing", 42, 68000, 11.3),
        ("Operations", 65, 58000, 6.1),
        ("HR", 18, 65000, 4.2),
        ("Finance", 12, 88000, 3.8),
    ]:
        ws2.append([dept, hc, sal, turn])

    # Sheet 3: Competitor analysis
    ws3 = wb.create_sheet("Competitors")
    ws3.append(["Company", "Market Share %", "Revenue ($M)", "Growth %", "Rating"])
    for comp, share, rev, growth, rating in [
        ("TechCorp", 32.1, 450, 18, "A+"),
        ("InnoSoft", 24.5, 320, 12, "A"),
        ("DataPrime", 18.8, 240, 25, "A-"),
        ("CloudFirst", 12.3, 160, 8, "B+"),
        ("Legacy Inc", 8.1, 105, -3, "B"),
    ]:
        ws3.append([comp, share, rev, growth, rating])

    wb.save(path)
    print(f"Created: {path}  (3 sheets)")


def create_simple_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = FIXTURES_DIR / "simple_1page.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4

    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, h - 72, "Q4 2025 Financial Summary")

    c.setFont("Helvetica", 11)
    y = h - 120
    for line in [
        "Quarterly Revenue: $12.8M (up 15% QoQ)",
        "Gross Margin: 68%",
        "Operating Income: $2.1M",
        "Free Cash Flow: $1.9M",
        "",
        "Key Metrics:",
        "- Cloud ARR: $8.2M",
        "- Net Revenue Retention: 118%",
        "- Payback Period: 14 months",
        "- LTV/CAC Ratio: 4.2x",
        "",
        "Recommendation: Strong buy. Cloud transition ahead of schedule.",
    ]:
        c.drawString(72, y, line)
        y -= 16

    c.save()
    print(f"Created: {path}")


if __name__ == "__main__":
    try:
        create_50page_pdf()
    except ImportError:
        print("SKIP: large_50page.pdf (reportlab not installed)")

    try:
        create_multisheet_xlsx()
    except ImportError:
        print("SKIP: multisheet.xlsx (openpyxl not installed)")

    try:
        create_simple_pdf()
    except ImportError:
        print("SKIP: simple_1page.pdf (reportlab not installed)")

    print("\nDone. Upload these via the frontend to verify.")
