"""Tests for common.attachments — text extraction from uploaded files."""

from pathlib import Path
import pytest

from common.attachments import extract_text, MAX_CHARS_PER_FILE


# ── Plain text extraction ──────────────────────────────────────────────


def test_extract_plain_text(tmp_path: Path):
    f = tmp_path / "note.txt"
    f.write_text("Hello world", encoding="utf-8")
    result = extract_text(f, "text/plain")
    assert result == "Hello world"


def test_extract_plain_text_empty(tmp_path: Path):
    f = tmp_path / "empty.txt"
    f.write_text("   ", encoding="utf-8")
    result = extract_text(f, "text/plain")
    assert result is None


def test_extract_plain_text_truncation(tmp_path: Path):
    f = tmp_path / "big.txt"
    content = "x" * (MAX_CHARS_PER_FILE + 1000)
    f.write_text(content, encoding="utf-8")
    result = extract_text(f, "text/plain")
    assert result is not None
    assert len(result) == MAX_CHARS_PER_FILE


# ── CSV extraction ─────────────────────────────────────────────────────


def test_extract_csv(tmp_path: Path):
    f = tmp_path / "data.csv"
    f.write_text("name,age\nAlice,30\nBob,25", encoding="utf-8")
    result = extract_text(f, "text/csv")
    assert result is not None
    assert "Alice" in result
    assert "Bob" in result


# ── PDF extraction ─────────────────────────────────────────────────────


def test_extract_pdf(tmp_path: Path):
    """Create a minimal PDF with pdfplumber-readable text."""
    pytest.importorskip("pdfplumber")
    rl_canvas = pytest.importorskip("reportlab.pdfgen.canvas")

    pdf_path = tmp_path / "test.pdf"
    c = rl_canvas.Canvas(str(pdf_path))
    c.drawString(72, 700, "This is a test PDF document.")
    c.save()

    result = extract_text(pdf_path, "application/pdf")
    assert result is not None
    assert "test PDF" in result


# ── XLSX extraction ────────────────────────────────────────────────────


def test_extract_xlsx(tmp_path: Path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    xlsx_path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Revenue"])
    ws.append(["Widget", 1200])
    ws.append(["Gadget", 3400])
    wb.save(xlsx_path)

    result = extract_text(xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert result is not None
    assert "Sales" in result
    assert "Widget" in result
    assert "1200" in result


# ── Multi-sheet XLSX ───────────────────────────────────────────────────


def test_extract_xlsx_multiple_sheets(tmp_path: Path):
    """Checklist item 3: Excel with 3 sheets → all appear in output."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    xlsx_path = tmp_path / "multi.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Revenue"
    ws1.append(["Quarter", "Amount"])
    ws1.append(["Q1", 1000])

    ws2 = wb.create_sheet("Costs")
    ws2.append(["Category", "Value"])
    ws2.append(["Salaries", 500])

    ws3 = wb.create_sheet("Summary")
    ws3.append(["Metric", "Result"])
    ws3.append(["Profit", 500])

    wb.save(xlsx_path)

    result = extract_text(xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert result is not None
    # All three sheet names should appear as markdown headers
    assert "Revenue" in result
    assert "Costs" in result
    assert "Summary" in result
    # Data from each sheet should be present
    assert "1000" in result
    assert "Salaries" in result
    assert "Profit" in result


# ── Unsupported mime ───────────────────────────────────────────────────


def test_extract_unsupported_mime(tmp_path: Path):
    f = tmp_path / "binary.exe"
    f.write_bytes(b"\x00\x01\x02")
    result = extract_text(f, "application/octet-stream")
    assert result is None


# ── Failure resilience ─────────────────────────────────────────────────


def test_extract_corrupt_file(tmp_path: Path):
    f = tmp_path / "corrupt.pdf"
    f.write_bytes(b"not a real pdf")
    result = extract_text(f, "application/pdf")
    assert result is None
