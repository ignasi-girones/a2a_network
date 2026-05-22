"""Best-effort text extraction from uploaded attachments.

Supports PDF (via pdfplumber), XLSX (via openpyxl), CSV, and plain text.
All extractors truncate output to avoid LLM context-window bombs.
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_CHARS_PER_FILE = 50_000
_XLSX_MAX_ROWS = 100
_XLSX_MAX_COLS = 20


def extract_text(path: Path, mime: str) -> str | None:
    """Return extracted text, or None on failure / unsupported mime."""
    try:
        if mime == "application/pdf":
            return _extract_pdf(path)
        if mime.endswith("spreadsheetml.sheet"):
            return _extract_xlsx(path)
        if mime in ("text/csv", "text/plain"):
            return _extract_plain(path)
    except Exception:
        logger.exception("Text extraction failed for %s (mime=%s)", path, mime)
    return None


def _extract_pdf(path: Path) -> str | None:
    import pdfplumber

    parts: list[str] = []
    total = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if not text:
                continue
            remaining = MAX_CHARS_PER_FILE - total
            if remaining <= 0:
                break
            parts.append(text[:remaining])
            total += len(parts[-1])
    return "\n\n".join(parts) if parts else None


def _extract_xlsx(path: Path) -> str | None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    total = 0
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(max_row=_XLSX_MAX_ROWS + 1, max_col=_XLSX_MAX_COLS, values_only=True))
        if not rows:
            continue
        header = "| " + " | ".join(str(c) if c is not None else "" for c in rows[0]) + " |"
        sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
        body_rows = [
            "| " + " | ".join(str(c) if c is not None else "" for c in row) + " |"
            for row in rows[1:_XLSX_MAX_ROWS + 1]
        ]
        table = f"### {sheet.title}\n\n{header}\n{sep}\n" + "\n".join(body_rows)
        remaining = MAX_CHARS_PER_FILE - total
        if remaining <= 0:
            break
        parts.append(table[:remaining])
        total += len(parts[-1])
    wb.close()
    return "\n\n".join(parts) if parts else None


def _extract_plain(path: Path) -> str | None:
    text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        return None
    return text[:MAX_CHARS_PER_FILE]
