"""Generate the A2A Debate Network presentation as .pptx"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette ──────────────────────────────────────────────────────────────
BLUE_600 = RGBColor(0x1B, 0x4F, 0x72)
BLUE_500 = RGBColor(0x2C, 0x6F, 0x9B)
BLUE_400 = RGBColor(0x51, 0x96, 0xB8)
BLUE_300 = RGBColor(0x7B, 0xAF, 0xCF)
BLUE_200 = RGBColor(0xB8, 0xD8, 0xEA)
BLUE_100 = RGBColor(0xE8, 0xF4, 0xFA)
BEIGE_100 = RGBColor(0xFB, 0xF7, 0xF2)
BEIGE_200 = RGBColor(0xF2, 0xE8, 0xD5)
BEIGE_300 = RGBColor(0xE3, 0xD3, 0xB8)
BEIGE_400 = RGBColor(0xD4, 0xB8, 0x96)
BEIGE_500 = RGBColor(0xC4, 0xA6, 0x7D)
DARK = RGBColor(0x2D, 0x34, 0x36)
MUTED = RGBColor(0x63, 0x6E, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x27, 0xAE, 0x60)
RED = RGBColor(0xC0, 0x39, 0x2B)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Helpers ──────────────────────────────────────────────────────────────

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, *,
             font_size=18, bold=False, color=DARK, alignment=PP_ALIGN.LEFT,
             font_name="Calibri", italic=False, line_spacing=1.3):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.font.italic = italic
    p.alignment = alignment
    p.space_after = Pt(0)
    p.line_spacing = Pt(font_size * line_spacing)
    return txBox


def add_rich_text(slide, left, top, width, height, runs, *,
                  alignment=PP_ALIGN.LEFT, line_spacing_pt=None):
    """runs = list of (text, font_size, bold, color, italic)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    for i, (txt, fs, bld, clr, ita) in enumerate(runs):
        run = p.add_run() if i > 0 else (p.runs[0] if p.runs else p.add_run())
        run.text = txt
        run.font.size = Pt(fs)
        run.font.bold = bld
        run.font.color.rgb = clr
        run.font.name = "Calibri"
        run.font.italic = ita
    if line_spacing_pt:
        p.line_spacing = Pt(line_spacing_pt)
    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color, *,
                     border_color=None, text="", font_size=14, font_color=DARK,
                     bold=False, alignment=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = alignment
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = font_color
        tf.paragraphs[0].font.name = "Calibri"
        tf.paragraphs[0].font.bold = bold
        shape.text_frame.margin_left = Pt(10)
        shape.text_frame.margin_right = Pt(10)
        shape.text_frame.margin_top = Pt(6)
        shape.text_frame.margin_bottom = Pt(6)
    return shape


def add_bullet_list(slide, left, top, width, height, items, *,
                    font_size=16, color=DARK, accent_color=BLUE_500,
                    spacing=Pt(8)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = spacing
        p.line_spacing = Pt(font_size * 1.4)
        # Bullet dot
        run = p.add_run()
        run.text = "  "
        run.font.size = Pt(font_size)
        run.font.color.rgb = accent_color
        run.font.name = "Calibri"
        # Text
        run2 = p.add_run()
        run2.text = item
        run2.font.size = Pt(font_size)
        run2.font.color.rgb = color
        run2.font.name = "Calibri"
    return txBox


def card_block(slide, left, top, width, height, title, body_lines, *,
               title_color=BLUE_500, bg_color=WHITE, border_color=BEIGE_300,
               font_size=13, title_size=16):
    add_rounded_rect(slide, left, top, width, height, bg_color,
                     border_color=border_color)
    add_text(slide, left + Inches(0.2), top + Inches(0.12),
             width - Inches(0.4), Inches(0.35), title,
             font_size=title_size, bold=True, color=title_color)
    y = top + Inches(0.5)
    for line in body_lines:
        add_text(slide, left + Inches(0.2), y,
                 width - Inches(0.4), Inches(0.3), line,
                 font_size=font_size, color=DARK)
        y += Inches(0.28)


def section_title(slide, text, *, color=BLUE_600):
    add_text(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.7),
             text, font_size=32, bold=True, color=color)


def subtitle_text(slide, text, top=Inches(1.15)):
    add_text(slide, Inches(0.8), top, Inches(11), Inches(0.4),
             text, font_size=17, color=MUTED)


# ── Build presentation ──────────────────────────────────────────────────

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]  # blank layout


# ═══════════════════════ SLIDE 1: TITLE ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
add_text(s, Inches(0.8), Inches(1.0), Inches(11), Inches(0.5),
         "UNIVERSITAT POLITECNICA DE CATALUNYA",
         font_size=13, color=BEIGE_500, bold=False)
add_text(s, Inches(0.8), Inches(1.8), Inches(11), Inches(1.2),
         "A2A Debate Network",
         font_size=52, bold=True, color=BLUE_600)
add_text(s, Inches(0.8), Inches(3.3), Inches(11), Inches(0.8),
         "Una red de agentes que deliberan\nusando el protocolo Agent-to-Agent",
         font_size=22, color=MUTED, line_spacing=1.5)
# Tags
tags = ["A2A Protocol", "Multi-Agent", "LLM", "Consensus"]
x = Inches(0.8)
for i, tag in enumerate(tags):
    clr = BLUE_200 if i % 2 == 0 else BEIGE_300
    tclr = BLUE_600 if i % 2 == 0 else DARK
    add_rounded_rect(s, x, Inches(4.6), Inches(1.8), Inches(0.4), clr,
                     text=tag, font_size=12, font_color=tclr, bold=True)
    x += Inches(2.0)
add_text(s, Inches(0.8), Inches(6.2), Inches(11), Inches(0.4),
         "Irina Regueiro  ·  2025", font_size=13, color=MUTED)


# ═══════════════════════ SLIDE 2: HOOK ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BLUE_600)
add_text(s, Inches(1.5), Inches(1.5), Inches(10), Inches(0.8),
         "Una pregunta", font_size=40, bold=True, color=BEIGE_200)
add_text(s, Inches(1.5), Inches(2.8), Inches(10), Inches(1.5),
         "¿Es mejor trabajar en remoto\no de manera presencial?",
         font_size=32, color=WHITE, line_spacing=1.5)
add_text(s, Inches(1.5), Inches(5.0), Inches(10), Inches(0.6),
         "Vamos a dejar que una red de agentes inteligentes debata esta cuestión.",
         font_size=18, color=BLUE_200)
# Note for presenter
add_text(s, Inches(1.5), Inches(6.2), Inches(10), Inches(0.4),
         "Personalizar con la pregunta del debate",
         font_size=11, color=BLUE_300, italic=True)


# ═══════════════════════ SLIDE 3: QUÉ ES A2A ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "¿Qué es el protocolo A2A?")
# Highlight box
add_rounded_rect(s, Inches(1.0), Inches(1.4), Inches(11.3), Inches(1.0),
                 BLUE_100, border_color=BLUE_200)
add_rich_text(s, Inches(1.3), Inches(1.5), Inches(10.7), Inches(0.8), [
    ("Un ", 19, False, DARK, False),
    ("estándar abierto", 19, True, BLUE_500, False),
    (" de Google que permite a agentes de IA ", 19, False, DARK, False),
    ("descubrirse", 19, True, BLUE_500, False),
    (", ", 19, False, DARK, False),
    ("comunicarse", 19, True, BLUE_500, False),
    (" y ", 19, False, DARK, False),
    ("colaborar", 19, True, BLUE_500, False),
    (" entre sí, independientemente de quién los haya creado.", 19, False, DARK, False),
], alignment=PP_ALIGN.CENTER, line_spacing_pt=28)

# 3 cards
cards = [
    ("Abierto", "Cualquier framework,\ncualquier lenguaje,\ncualquier proveedor"),
    ("Interoperable", "Agentes de distintos\nequipos colaboran\nsin adaptadores"),
    ("Descubrible", "Los agentes publican sus\ncapacidades y se\nencuentran solos"),
]
for i, (title, body) in enumerate(cards):
    x = Inches(1.0) + Inches(i * 3.9)
    add_rounded_rect(s, x, Inches(3.0), Inches(3.5), Inches(3.0), WHITE,
                     border_color=BEIGE_300)
    add_text(s, x + Inches(0.3), Inches(3.3), Inches(2.9), Inches(0.4),
             title, font_size=20, bold=True, color=BLUE_500)
    add_text(s, x + Inches(0.3), Inches(3.9), Inches(2.9), Inches(1.8),
             body, font_size=15, color=DARK, line_spacing=1.5)


# ═══════════════════════ SLIDE 4: AGENT CARD ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Agent Card")
subtitle_text(s, "La identidad pública de cada agente")

# JSON block
json_text = """{
  "name": "Specialized Agent (AE1)",
  "version": "1.0.0",
  "capabilities": {
    "streaming": true
  },
  "skills": [{
    "id": "debate",
    "name": "Debate",
    "description": "Argues a topic
      from an assigned role...",
    "tags": ["deliberative",
             "multi-agent"]
  }]
}"""
add_rounded_rect(s, Inches(0.8), Inches(1.7), Inches(5.5), Inches(5.0),
                 BLUE_600)
add_text(s, Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.7),
         json_text, font_size=14, color=BLUE_100, font_name="Consolas",
         line_spacing=1.3)

# Right side cards
right_cards = [
    ("Metadata", "Nombre, versión, proveedor, URL"),
    ("Capabilities", "Streaming, push notifications,\nmodo de input/output"),
    ("Skills", "Qué sabe hacer, cuándo usarlo,\nqué formato espera"),
]
for i, (title, body) in enumerate(right_cards):
    y = Inches(1.7) + Inches(i * 1.55)
    card_block(s, Inches(6.8), y, Inches(5.5), Inches(1.35),
               title, [body], title_size=18, font_size=14)

add_text(s, Inches(0.8), Inches(6.9), Inches(11), Inches(0.4),
         "Servida en /.well-known/agent-card.json — el punto de entrada del descubrimiento",
         font_size=12, color=MUTED, alignment=PP_ALIGN.CENTER)


# ═══════════════════════ SLIDE 5: DISCOVERY ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Agent Discovery")
subtitle_text(s, "Los agentes se encuentran y se registran automáticamente")

# Flow 1
boxes1 = [
    ("Agente se inicia", BLUE_200, BLUE_600),
    ("Publica Agent Card", BEIGE_300, DARK),
    ("Se registra en\nel Orquestador", BLUE_200, BLUE_600),
]
for i, (txt, bg, fc) in enumerate(boxes1):
    x = Inches(1.0) + Inches(i * 4.2)
    add_rounded_rect(s, x, Inches(2.0), Inches(3.2), Inches(1.0), bg,
                     text=txt, font_size=16, font_color=fc, bold=True)
    if i < 2:
        add_text(s, x + Inches(3.3), Inches(2.15), Inches(0.8), Inches(0.6),
                 "→", font_size=30, color=BLUE_400, alignment=PP_ALIGN.CENTER)

# Flow 2
boxes2 = [
    ("Orquestador\nconsulta Registry", WHITE, DARK),
    ("Lee skills de\ncada agente", BEIGE_300, DARK),
    ("Planifica\nel debate", BLUE_200, BLUE_600),
]
for i, (txt, bg, fc) in enumerate(boxes2):
    x = Inches(1.0) + Inches(i * 4.2)
    add_rounded_rect(s, x, Inches(3.8), Inches(3.2), Inches(1.0), bg,
                     text=txt, font_size=16, font_color=fc, bold=True,
                     border_color=BEIGE_300 if bg == WHITE else None)
    if i < 2:
        add_text(s, x + Inches(3.3), Inches(3.95), Inches(0.8), Inches(0.6),
                 "→", font_size=30, color=BLUE_400, alignment=PP_ALIGN.CENTER)

# Bottom card
add_rounded_rect(s, Inches(2.5), Inches(5.5), Inches(8.3), Inches(1.0),
                 WHITE, border_color=BEIGE_300)
add_rich_text(s, Inches(2.8), Inches(5.65), Inches(7.7), Inches(0.7), [
    ("El orquestador ", 17, False, DARK, False),
    ("no sabe de antemano", 17, True, BLUE_500, False),
    (" qué agentes existen.\nLee sus skills y decide cómo usarlos.", 17, False, DARK, False),
], alignment=PP_ALIGN.CENTER, line_spacing_pt=26)


# ═══════════════════════ SLIDE 6: COMUNICACIÓN ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Comunicación A2A")

card_block(s, Inches(0.8), Inches(1.4), Inches(5.5), Inches(2.5),
           "JSON-RPC 2.0", [
               "Protocolo de transporte estándar",
               "",
               "· SendMessage — síncrono",
               "· SendStreamingMessage — SSE",
           ], font_size=15, title_size=20)

card_block(s, Inches(7.0), Inches(1.4), Inches(5.5), Inches(2.5),
           "Task Lifecycle", [
               "Cada petición crea un Task",
               "con estados:",
               "",
               "submitted → working → completed",
           ], font_size=15, title_size=20)

add_rounded_rect(s, Inches(2.0), Inches(4.5), Inches(9.3), Inches(1.2),
                 WHITE, border_color=BEIGE_300)
add_rich_text(s, Inches(2.3), Inches(4.7), Inches(8.7), Inches(0.8), [
    ("El orquestador usa ", 17, False, DARK, False),
    ("streaming", 17, True, BLUE_500, False),
    (" hacia el frontend (progreso en tiempo real)\ny llamadas ", 17, False, DARK, False),
    ("síncronas", 17, True, BLUE_500, False),
    (" a los workers (espera respuesta completa).", 17, False, DARK, False),
], alignment=PP_ALIGN.CENTER, line_spacing_pt=26)


# ═══════════════════════ SLIDE 7: CASO DE USO ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_200)
section_title(s, "Nuestro caso de uso")

add_rich_text(s, Inches(0.8), Inches(1.2), Inches(11), Inches(0.6), [
    ("Una ", 20, False, DARK, False),
    ("deliberación multi-agente", 20, True, BLUE_500, False),
    (" donde agentes con distintos roles debaten un tema y buscan consenso.", 20, False, DARK, False),
], line_spacing_pt=30)

card_block(s, Inches(0.8), Inches(2.3), Inches(5.5), Inches(4.0),
           "¿Por qué debates?", [
               "· Explota la naturaleza distribuida de A2A",
               "",
               "· Cada agente es independiente",
               "  con su propio LLM",
               "",
               "· El consenso emerge de la",
               "  interacción, no de un solo modelo",
           ], font_size=15, title_size=20, bg_color=WHITE)

card_block(s, Inches(7.0), Inches(2.3), Inches(5.5), Inches(4.0),
           "¿Qué lo hace interesante?", [
               "· Distintos modelos =",
               "  distintas perspectivas",
               "",
               "· El orquestador no sabe la respuesta",
               "",
               "· Agentes externos pueden",
               "  unirse en vivo",
           ], font_size=15, title_size=20, bg_color=WHITE)


# ═══════════════════════ SLIDE 8: ARQUITECTURA ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Arquitectura")

# Orchestrator (top bar)
add_rounded_rect(s, Inches(1.5), Inches(1.4), Inches(10.3), Inches(1.0),
                 BLUE_100, border_color=BLUE_200,
                 text="Orquestador  —  Gemini 2.5 Flash  —  planifica, coordina, evalúa consenso",
                 font_size=17, font_color=BLUE_600, bold=True)

# Agent cards row
agents = [
    ("Normalizer", "Gemini Flash", "preproceso", BEIGE_300),
    ("AE1", "Mistral Large", "debate", BLUE_200),
    ("AE2", "Cerebras Llama", "debate", BLUE_200),
    ("AE3", "Groq Llama", "debate", BLUE_200),
]
for i, (name, model, skill, clr) in enumerate(agents):
    x = Inches(1.0) + Inches(i * 3.1)
    add_rounded_rect(s, x, Inches(2.9), Inches(2.8), Inches(1.8),
                     WHITE, border_color=BEIGE_300)
    add_text(s, x + Inches(0.2), Inches(3.1), Inches(2.4), Inches(0.4),
             name, font_size=18, bold=True, color=BLUE_500, alignment=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(3.55), Inches(2.4), Inches(0.3),
             model, font_size=12, color=MUTED, alignment=PP_ALIGN.CENTER)
    add_rounded_rect(s, x + Inches(0.7), Inches(4.0), Inches(1.4), Inches(0.35),
                     clr, text=skill, font_size=11, font_color=BLUE_600 if clr == BLUE_200 else DARK, bold=True)

# Bottom row
bottom = [
    ("Feedback", "veredicto"),
    ("MCP Tools", "herramientas"),
]
for i, (name, skill) in enumerate(bottom):
    x = Inches(3.5) + Inches(i * 3.5)
    add_rounded_rect(s, x, Inches(5.2), Inches(3.0), Inches(1.0),
                     WHITE, border_color=BEIGE_300)
    add_text(s, x + Inches(0.2), Inches(5.35), Inches(2.6), Inches(0.3),
             name, font_size=17, bold=True, color=BLUE_500, alignment=PP_ALIGN.CENTER)
    add_rounded_rect(s, x + Inches(0.8), Inches(5.75), Inches(1.4), Inches(0.3),
                     BEIGE_300, text=skill, font_size=10, font_color=DARK, bold=True)

add_text(s, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4),
         "Cada agente es un proceso independiente con su propio LLM, conectado via A2A",
         font_size=12, color=MUTED, alignment=PP_ALIGN.CENTER)


# ═══════════════════════ SLIDE 9: FLUJO ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Flujo de ejecución")
subtitle_text(s, "7 fases — el LLM decide el plan, no el código")

phases = ["Discover", "Plan", "Capacity", "Execute", "Consensus", "Feedback", "Synthesize"]
for i, phase in enumerate(phases):
    x = Inches(0.5) + Inches(i * 1.75)
    bg = BLUE_100 if i == 4 else WHITE
    bc = BLUE_200 if i == 4 else BEIGE_300
    add_rounded_rect(s, x, Inches(2.0), Inches(1.5), Inches(1.0), bg,
                     border_color=bc)
    add_text(s, x, Inches(2.1), Inches(1.5), Inches(0.4),
             str(i+1), font_size=24, bold=True, color=BLUE_500, alignment=PP_ALIGN.CENTER)
    add_text(s, x, Inches(2.55), Inches(1.5), Inches(0.3),
             phase, font_size=12, color=DARK, alignment=PP_ALIGN.CENTER)
    if i < 6:
        add_text(s, x + Inches(1.5), Inches(2.15), Inches(0.3), Inches(0.6),
                 "→", font_size=20, color=BLUE_400, alignment=PP_ALIGN.CENTER)

card_block(s, Inches(0.8), Inches(3.5), Inches(5.5), Inches(3.2),
           "Lo que decide el LLM", [
               "· Qué skills usar y en qué orden",
               "· Qué roles asignar a cada agente",
               "· Cuándo extender el debate",
               "· Si añadir agentes nuevos",
           ], font_size=15, title_size=18)

card_block(s, Inches(7.0), Inches(3.5), Inches(5.5), Inches(3.2),
           "Lo que está hardcodeado", [
               "· Métricas de consenso (embeddings)",
               "· Umbral de convergencia (0.75)",
               "· Máximo de extensiones (3)",
               "· Pesos de los componentes del score",
           ], font_size=15, title_size=18)


# ═══════════════════════ SLIDE 10: CONSENSO ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Consenso empírico")
subtitle_text(s, "Sin opinión del LLM — solo matemáticas")

card_block(s, Inches(0.8), Inches(1.7), Inches(5.8), Inches(2.0),
           "Eje AE1 ↔ AE2", [
               "Las aperturas de AE1 y AE2 definen",
               "posición 0.0 y 1.0. Cada texto posterior",
               "se ubica por similitud coseno",
               "con esos anclas.",
           ], font_size=15, title_size=18)

card_block(s, Inches(0.8), Inches(4.1), Inches(5.8), Inches(2.6),
           "4 componentes", [
               "· Dispersión (40%) — qué tan cerca están",
               "· Similitud (35%) — semántica entre textos",
               "· Movimiento (15%) — ¿se están moviendo?",
               "· Concesiones (10%) — \"tienes razón\"",
           ], font_size=15, title_size=18)

# Formula card
add_rounded_rect(s, Inches(7.0), Inches(1.7), Inches(5.5), Inches(5.0),
                 BLUE_100, border_color=BLUE_200)
add_text(s, Inches(7.3), Inches(2.0), Inches(5.0), Inches(0.4),
         "Fórmula del score", font_size=14, color=MUTED, alignment=PP_ALIGN.CENTER)
add_text(s, Inches(7.3), Inches(2.8), Inches(5.0), Inches(2.5),
         "score = 0.40 · dispersión\n"
         "      + 0.35 · similitud\n"
         "      + 0.15 · movimiento\n"
         "      + 0.10 · concesiones",
         font_size=20, color=BLUE_600, font_name="Consolas", line_spacing=1.8,
         alignment=PP_ALIGN.CENTER)
add_rich_text(s, Inches(7.3), Inches(5.3), Inches(5.0), Inches(0.6), [
    ("Si score ≥ ", 18, False, DARK, False),
    ("0.75", 28, True, BLUE_500, False),
    ("  →  consenso", 18, False, DARK, False),
], alignment=PP_ALIGN.CENTER)


# ═══════════════════════ SLIDE 11: DYNAMIC AGENTS ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BLUE_600)
add_text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.7),
         "Adición dinámica de agentes", font_size=32, bold=True, color=BEIGE_200)
add_text(s, Inches(0.8), Inches(1.2), Inches(11), Inches(0.4),
         "La potencia real del protocolo A2A aplicada a nuestro caso de uso",
         font_size=17, color=BLUE_200)

# Round 1
add_rounded_rect(s, Inches(0.8), Inches(2.0), Inches(5.5), Inches(3.0),
                 RGBColor(0x25, 0x5E, 0x85), border_color=RGBColor(0x3A, 0x7A, 0xA8))
add_text(s, Inches(1.1), Inches(2.2), Inches(5.0), Inches(0.4),
         "Ronda 1", font_size=20, bold=True, color=BEIGE_200)
add_text(s, Inches(1.1), Inches(2.8), Inches(5.0), Inches(2.0),
         "AE1, AE2 y AE3 debaten.\nNo llegan a consenso.\n\n"
         "El orquestador detecta agentes\ndisponibles en el registry que\nno están participando.",
         font_size=15, color=BLUE_200, line_spacing=1.5)

# Round 2
add_rounded_rect(s, Inches(7.0), Inches(2.0), Inches(5.5), Inches(3.0),
                 RGBColor(0x25, 0x5E, 0x85), border_color=RGBColor(0x3A, 0x7A, 0xA8))
add_text(s, Inches(7.3), Inches(2.2), Inches(5.0), Inches(0.4),
         "Ronda 2", font_size=20, bold=True, color=BEIGE_200)
add_text(s, Inches(7.3), Inches(2.8), Inches(5.0), Inches(2.0),
         "El planner decide: \"Un experto\nlegal ayudaría a desbloquear.\"\n\n"
         "AE4 se incorpora al debate\ncon contexto completo de\nlo discutido.",
         font_size=15, color=BLUE_200, line_spacing=1.5)

add_text(s, Inches(1.0), Inches(5.5), Inches(11.3), Inches(1.2),
         "El agente añadido puede ser de otro equipo, otra organización, otro país.\n"
         "Solo necesita publicar una Agent Card con skill \"debate\".",
         font_size=18, color=WHITE, alignment=PP_ALIGN.CENTER, line_spacing=1.6)


# ═══════════════════════ SLIDE 12: CÓMO FUNCIONA ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "¿Cómo funciona?")

flow = [
    ("Score < 0.75", BEIGE_300, DARK),
    ("Buscar agentes\ndebate no\nparticipando", BLUE_200, BLUE_600),
    ("Planner evalúa\nsi añadirlos", BEIGE_300, DARK),
    ("Extensión con\nN+1 agentes", BLUE_200, BLUE_600),
]
for i, (txt, bg, fc) in enumerate(flow):
    x = Inches(0.6) + Inches(i * 3.3)
    add_rounded_rect(s, x, Inches(1.5), Inches(2.8), Inches(1.2), bg,
                     text=txt, font_size=15, font_color=fc, bold=True)
    if i < 3:
        add_text(s, x + Inches(2.85), Inches(1.65), Inches(0.5), Inches(0.6),
                 "→", font_size=28, color=BLUE_400, alignment=PP_ALIGN.CENTER)

card_block(s, Inches(0.6), Inches(3.3), Inches(3.6), Inches(2.5),
           "Filtrado inteligente", [
               "Solo agentes con skill debate.",
               "Nunca normalizer, feedback",
               "ni herramientas.",
           ], font_size=14, title_size=17)

card_block(s, Inches(4.8), Inches(3.3), Inches(3.6), Inches(2.5),
           "Decisión del LLM", [
               "El planner decide si la",
               "expertise del nuevo agente",
               "ayuda al deadlock actual.",
           ], font_size=14, title_size=17)

card_block(s, Inches(9.0), Inches(3.3), Inches(3.6), Inches(2.5),
           "Catch-up automático", [
               "El nuevo agente recibe",
               "todo el contexto previo",
               "via depends_on.",
           ], font_size=14, title_size=17)


# ═══════════════════════ SLIDE 13: DEPLOY ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Despliegue y monitorización")

card_block(s, Inches(0.8), Inches(1.4), Inches(5.5), Inches(2.4),
           "Infraestructura", [
               "· Docker Compose con 8 servicios",
               "· mTLS entre agentes en producción",
               "· Healthchecks con dependencias",
               "· VM en la UPC (nattech.fib.upc.edu)",
           ], font_size=14, title_size=18)

card_block(s, Inches(0.8), Inches(4.2), Inches(5.5), Inches(2.2),
           "Persistencia", [
               "· SQLite para historial de debates",
               "· Replay de SSE events en F5",
               "· Volúmenes Docker para datos",
           ], font_size=14, title_size=18)

card_block(s, Inches(7.0), Inches(1.4), Inches(5.5), Inches(2.4),
           "Observabilidad", [
               "· Prometheus — métricas de cada agente",
               "· Loki + Promtail — logs centralizados",
               "· Grafana — dashboards en tiempo real",
           ], font_size=14, title_size=18, bg_color=BLUE_100, border_color=BLUE_200,
           title_color=BLUE_600)

card_block(s, Inches(7.0), Inches(4.2), Inches(5.5), Inches(2.2),
           "¿Qué monitorizamos?", [
               "· Latencia LLM por agente y modelo",
               "· Tasa de error de cada worker",
               "· Correlación de logs por debate_id",
               "· Heatmap de deliberación",
           ], font_size=14, title_size=18)


# ═══════════════════════ SLIDE 14: LANZAR DEMO ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_200)
add_text(s, Inches(1.5), Inches(1.8), Inches(10), Inches(0.4),
         "MOMENTO DE LA VERDAD", font_size=14, color=MUTED)
add_text(s, Inches(1.5), Inches(2.5), Inches(10), Inches(1.0),
         "Lanzamos la demo", font_size=44, bold=True, color=BLUE_600)
add_text(s, Inches(1.5), Inches(4.0), Inches(10), Inches(0.6),
         "Mientras los agentes debaten, hablemos del futuro del protocolo.",
         font_size=20, color=DARK)


# ═══════════════════════ SLIDE 15: TESIS ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)

# Quote
add_rounded_rect(s, Inches(0.6), Inches(0.6), Inches(0.12), Inches(1.8),
                 BEIGE_400)
add_text(s, Inches(1.1), Inches(0.6), Inches(11), Inches(1.8),
         "A2A será al ecosistema de agentes\n"
         "lo que HTTP fue a la web:\n"
         "el protocolo que permite que agentes\n"
         "de distintos proveedores colaboren\n"
         "sin conocerse previamente.",
         font_size=24, color=BLUE_600, italic=True, line_spacing=1.4)

add_text(s, Inches(0.8), Inches(3.0), Inches(11), Inches(0.5),
         "¿Por qué creemos esto?", font_size=22, bold=True, color=BLUE_500,
         alignment=PP_ALIGN.CENTER)

cards_data = [
    ("El problema es real", "Hoy cada framework de\nagentes es un silo.\nLangChain no habla con\nCrewAI ni con AutoGen."),
    ("Respaldo de la industria", "Google, Salesforce, SAP,\nAtlassian, MongoDB y\nmás de 50 empresas."),
    ("Diseño pragmático", "JSON-RPC, HTTP, Agent Cards.\nUsa estándares probados,\nno reinventa la rueda."),
]
for i, (title, body) in enumerate(cards_data):
    x = Inches(0.6) + Inches(i * 4.2)
    card_block(s, x, Inches(3.7), Inches(3.8), Inches(2.8),
               title, body.split("\n"), font_size=15, title_size=17)


# ═══════════════════════ SLIDE 16: DIFERENCIAL ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Donde A2A puede ser diferencial")

diff_cards = [
    ("Empresa", "Un agente de RRHH de SAP\nnegocia vacaciones con un\nagente de planning de Jira."),
    ("Marketplaces", "Publicar agentes especializados\ncomo servicios. Descubrimiento\nautomático por skills."),
    ("Cadenas de valor", "Agente de un proveedor coordina\ncon agente de logística de\notra empresa en tiempo real."),
    ("Investigación", "Agentes de distintos laboratorios\nanalizan un mismo dataset desde\nperspectivas diferentes."),
]
for i, (title, body) in enumerate(diff_cards):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + Inches(col * 6.2)
    y = Inches(1.5) + Inches(row * 2.7)
    card_block(s, x, y, Inches(5.8), Inches(2.3),
               title, body.split("\n"), font_size=15, title_size=18)


# ═══════════════════════ SLIDE 17: RETOS ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "Retos y preguntas abiertas")

retos = [
    ("Confianza y seguridad", "¿Cómo confías en un agente que\nno conoces? Autenticación,\nautorización, sandboxing.", RED),
    ("Observabilidad a escala", "Con decenas de agentes, el\ndebugging se vuelve exponencial.\nTracing distribuido nativo.", RED),
    ("Competencia con MCP", "Model Context Protocol (Anthropic)\nresuelve parte del problema.\n¿Convivencia o guerra?", RED),
    ("Nuestra perspectiva", "MCP conecta modelos con\nherramientas. A2A conecta\nagentes entre sí.\nSon complementarios.", GREEN),
]
for i, (title, body, tcolor) in enumerate(retos):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + Inches(col * 6.2)
    y = Inches(1.5) + Inches(row * 2.7)
    card_block(s, x, y, Inches(5.8), Inches(2.3),
               title, body.split("\n"), font_size=15, title_size=18,
               title_color=tcolor)


# ═══════════════════════ SLIDE 18: A2A + MCP ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
section_title(s, "A2A + MCP = el stack completo")
subtitle_text(s, "No compiten — cada uno resuelve una capa distinta")

# Table header
headers = ["", "MCP (Anthropic)", "A2A (Google)"]
col_widths = [Inches(2.5), Inches(4.0), Inches(4.0)]
x_start = Inches(1.2)
for i, (h, w) in enumerate(zip(headers, col_widths)):
    x = x_start + sum(cw for cw in [Inches(0)] + list(col_widths[:i]))
    add_rounded_rect(s, x, Inches(1.8), w, Inches(0.55),
                     BLUE_500 if i > 0 or not h else BLUE_500,
                     text=h, font_size=15, font_color=WHITE, bold=True)

rows_data = [
    ("Conecta", "Modelo ↔ Herramientas", "Agente ↔ Agente"),
    ("Analogía", "USB para un dispositivo", "HTTP entre servidores"),
    ("Ejemplo", "Agente accede a base de datos", "Dos agentes negocian un contrato"),
    ("Scope", "Dentro de un agente", "Entre agentes independientes"),
]
for r, (label, mcp, a2a) in enumerate(rows_data):
    y = Inches(2.4) + Inches(r * 0.6)
    bg = BEIGE_100 if r % 2 == 0 else WHITE
    vals = [label, mcp, a2a]
    x = x_start
    for i, (v, w) in enumerate(zip(vals, col_widths)):
        add_rounded_rect(s, x, y, w, Inches(0.55), bg,
                         text=v, font_size=14,
                         font_color=DARK if i > 0 else BLUE_600,
                         bold=(i == 0),
                         alignment=PP_ALIGN.LEFT)
        x += w

add_rounded_rect(s, Inches(2.5), Inches(5.2), Inches(8.3), Inches(1.2),
                 WHITE, border_color=BEIGE_300)
add_rich_text(s, Inches(2.8), Inches(5.35), Inches(7.7), Inches(0.8), [
    ("En nuestro proyecto ", 17, False, DARK, False),
    ("usamos ambos", 17, True, BLUE_500, False),
    (": MCP para que los agentes accedan a herramientas,\nA2A para que se comuniquen entre ellos.", 17, False, DARK, False),
], alignment=PP_ALIGN.CENTER, line_spacing_pt=26)


# ═══════════════════════ SLIDE 19: LECCIONES ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_200)
section_title(s, "Lo que hemos aprendido")

lessons = [
    ("El consenso es difícil", "Los LLMs tienden al\n\"ambos tienen razón\".\nForzar convergencia real\nrequiere prompt engineering\nagresivo."),
    ("Rate limits everywhere", "Con 4+ agentes en paralelo,\nlos TPM de los proveedores\nse agotan rápido.\nLa diversidad de modelos\nayuda."),
    ("Métricas > opinión", "Medir consenso con\nembeddings es más fiable\nque pedirle al LLM\nque se autoevalúe."),
]
for i, (title, body) in enumerate(lessons):
    x = Inches(0.6) + Inches(i * 4.2)
    card_block(s, x, Inches(1.5), Inches(3.8), Inches(4.2),
               title, body.split("\n"), font_size=15, title_size=18,
               bg_color=WHITE)


# ═══════════════════════ SLIDE 20: DEMO RESULTS ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BEIGE_100)
add_text(s, Inches(1.5), Inches(1.8), Inches(10), Inches(0.4),
         "Y AHORA...", font_size=14, color=MUTED)
add_text(s, Inches(1.5), Inches(2.5), Inches(10), Inches(1.0),
         "¿Qué han decidido\nlos agentes?",
         font_size=42, bold=True, color=BLUE_600, line_spacing=1.3)
add_text(s, Inches(1.5), Inches(4.5), Inches(10), Inches(0.6),
         "Veamos el resultado del debate en vivo.",
         font_size=20, color=DARK)


# ═══════════════════════ SLIDE 21: CIERRE ═══════════════════════
s = prs.slides.add_slide(blank)
set_slide_bg(s, BLUE_600)
add_text(s, Inches(1.5), Inches(2.0), Inches(10), Inches(1.0),
         "Gracias", font_size=52, bold=True, color=BEIGE_200,
         alignment=PP_ALIGN.CENTER)

tags = ["Agent-to-Agent Protocol v1.0", "Multi-Agent Deliberation", "Dynamic Agent Discovery"]
x = Inches(2.0)
for i, tag in enumerate(tags):
    w = Inches(3.2) if i == 0 else Inches(3.0)
    add_rounded_rect(s, x, Inches(3.5), w, Inches(0.45),
                     BEIGE_300 if i != 1 else RGBColor(0x25, 0x5E, 0x85),
                     text=tag, font_size=12,
                     font_color=DARK if i != 1 else BEIGE_200, bold=True)
    x += w + Inches(0.15)

add_text(s, Inches(1.5), Inches(4.5), Inches(10), Inches(0.5),
         "¿Preguntas?", font_size=22, color=BLUE_200,
         alignment=PP_ALIGN.CENTER)
add_text(s, Inches(1.5), Inches(6.0), Inches(10), Inches(0.4),
         "Irina Regueiro  ·  UPC  ·  2025", font_size=13, color=BLUE_300,
         alignment=PP_ALIGN.CENTER)


# ── Save ─────────────────────────────────────────────────────────────
out = "/home/irina/proyectos/a2a_network/docs/A2A_Debate_Network.pptx"
prs.save(out)
print(f"Saved to {out}")
